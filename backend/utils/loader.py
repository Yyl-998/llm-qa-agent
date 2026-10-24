"""文档加载与切分：支持 pdf / txt / md / docx。"""
from pathlib import Path

from langchain_core.documents import Document

from backend.config import DOCS_DIR


def load_document(path: Path) -> list[Document]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from langchain_community.document_loaders import PyPDFLoader

        docs = PyPDFLoader(str(path)).load()
    elif ext == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        docs = [Document(page_content=text)]
    else:  # .txt / .md
        docs = [Document(page_content=_read_text(path))]
    for d in docs:
        d.metadata["source"] = path.name
    return docs


def split_documents(docs: list[Document]) -> list[Document]:
    """中文文档切分：片段较小以适配 bge 512 token 窗口。"""
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(chunk_size=400, chunk_overlap=60)
    return splitter.split_documents(docs)


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")
